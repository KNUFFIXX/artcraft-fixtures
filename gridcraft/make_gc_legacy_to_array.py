#!/usr/bin/env python3
"""make_gc_legacy_to_array.py (DRACO-AC03) - legacy (pre-dynamic-array) formulas over whole-column names, as Excel 2007-2016
and openpyxl write them (plain <f>, no cm): Gesamtkosten = Calc!$D:$D, Gesamt_EH = Calc!$F:$F, Tage = a row of day counts.
H2:AA21 (20x20) = IFERROR(Gesamtkosten/Gesamt_EH*$G2*Tage/365,0): in Excel each cell intersects implicitly with its own row
and column (one number). Saved as an array formula, each cell evaluates 1,048,576 x 20 values instead."""
import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as L

def make(path, nrows, ncols):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Calc'
    ws['A1'] = 'Days per period'
    for j in range(ncols): ws.cell(row=1, column=8 + j, value=30 + j)
    for i in range(nrows):
        r = 2 + i
        ws.cell(row=r, column=4, value=1000 + 10 * i)
        ws.cell(row=r, column=6, value=100 + i)
        ws.cell(row=r, column=7, value=5 + i % 7)
        for j in range(ncols):
            ws.cell(row=r, column=8 + j, value=f'=IFERROR(Gesamtkosten/Gesamt_EH*$G{r}*Tage/365,0)')
    last = L(8 + ncols - 1)
    for n, ref in (('Gesamtkosten', 'Calc!$D:$D'), ('Gesamt_EH', 'Calc!$F:$F'), ('Tage', f'Calc!$H$1:${last}$1')):
        wb.defined_names[n] = DefinedName(n, attr_text=ref)
    wb.save(path)

if __name__ == '__main__':
    make('gc_legacy_to_array_small.xlsx', 5, 5)
    make('gc_legacy_to_array.xlsx', 20, 20)
