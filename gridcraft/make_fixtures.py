#!/usr/bin/env python3
"""make_fixtures.py - builds every GridCraft fixture in this folder (Python 3 + openpyxl 3.1). All content is synthetic.

  gc_table_thisrow.xlsx          table with calculated columns stored as Excel stores them (Sales[[#This Row],[Qty]]), one column name with a space
  gc_table_header_newline.xlsx   table whose header cells contain a line break; column names carry Excel's _x000a_ escape
  gc_normal_font_arial10.xlsx    hand-written package whose Normal style (cellXfs 0) is Arial 10
  gc_normal_font_aptos.xlsx      same with Aptos Narrow 11
  gc_implicit_intersection.xlsx  legacy formulas (no dynamic-array metadata) over whole-column names and ranges
  gc_scoped_names.xlsx           global names plus sheet-scoped duplicates
  gc_coerce_de.xlsx              text operands with German separators; expected values of a de-DE Excel in D1
  gc_sort_collation.xlsx         SORT/SORTBY over strings with hyphens, underscores, case and umlauts; Excel's order in G1
"""
import os, zipfile, openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn, TableFormula
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
def p(name): return os.path.join(HERE, name)

def table(ws, name, ref, headers, calc=None):
    t = Table(displayName=name, ref=ref); t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    cols = []
    for i, h in enumerate(headers, start=1):
        c = TableColumn(id=i, name=h)
        if calc and h in calc: c.calculatedColumnFormula = TableFormula(attr_text=calc[h])
        cols.append(c)
    t.tableColumns = cols; ws.add_table(t); return t

def patch_parts(path, prefix, repl):
    """byte replacements inside the zip parts whose name starts with prefix"""
    tmp = path + '.tmp'; src = zipfile.ZipFile(path); out = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for n in src.namelist():
        d = src.read(n)
        if n.startswith(prefix):
            for a, b in repl: d = d.replace(a, b)
        out.writestr(n, d)
    out.close(); src.close(); os.replace(tmp, path)

# --- gc_table_thisrow.xlsx ---------------------------------------------------------------------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sales'
hdr = ['Item', 'Qty', 'Unit Price', 'Total', 'Share']; ws.append(hdr)
for i in range(1, 6):
    ws.append(['it%d' % i, i, 2.5 * i, '=Sales[[#This Row],[Qty]]*Sales[[#This Row],[Unit Price]]', '=Sales[[#This Row],[Total]]/SUM(Sales[Total])'])
table(ws, 'Sales', 'A1:E6', hdr, {'Total': 'Sales[[#This Row],[Qty]]*Sales[[#This Row],[Unit Price]]', 'Share': 'Sales[[#This Row],[Total]]/SUM(Sales[Total])'})
ws['G1'] = '=SUM(Sales[Total])'; ws['G2'] = 'Expected: Excel opens the file after a GridCraft save; D2 = 2.5'
wb.save(p('gc_table_thisrow.xlsx'))

# --- gc_table_header_newline.xlsx --------------------------------------------------------------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sheet1'
hdr = ['Art der \nZahlung', 'Kalt-\nmiete', 'Betrag']; ws.append(hdr)
ws.append([1, 2, '=A2+B2']); ws.append([3, 4, '=A3+B3'])
table(ws, 'Zahlungen', 'A1:C3', hdr)   # a table name must not look like a cell reference (TBL1 would be cell TBL1)
ws['E1'] = 'Expected after a GridCraft save: A1 still "Art der \\nZahlung" (LEN = 16), B1 "Kalt-\\nmiete" (LEN = 11)'
ws['E2'] = '=LEN(A1)'; ws['E3'] = '=LEN(B1)'
wb.save(p('gc_table_header_newline.xlsx'))
patch_parts(p('gc_table_header_newline.xlsx'), 'xl/tables/', [(b'&#10;', b'_x000a_')])   # Excel's spelling of the line break

# --- gc_normal_font_*.xlsx (hand-written package) --------------------------------------------
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
def make_min_xlsx(path, rows_xml, font0):
    ct = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
          '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
          '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    wbrels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
              '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
              '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>')
    wbx = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
           '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>' % NS)
    styles = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><styleSheet xmlns="%s">'
              '<fonts count="2"><font><sz val="%s"/><name val="%s"/></font><font><b/><sz val="%s"/><name val="%s"/></font></fonts>'
              '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
              '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
              '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
              '<cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/></cellXfs>'
              '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>') % (NS, font0[1], font0[0], font0[1], font0[0])
    sheet = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="%s"><sheetData>%s</sheetData></worksheet>' % (NS, ''.join(rows_xml)))
    z = zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED)
    for n, d in (('[Content_Types].xml', ct), ('_rels/.rels', rels), ('xl/workbook.xml', wbx), ('xl/_rels/workbook.xml.rels', wbrels), ('xl/styles.xml', styles), ('xl/worksheets/sheet1.xml', sheet)):
        z.writestr(n, d)
    z.close()
rows = ['<row r="1"><c r="A1" t="inlineStr"><is><t>default font cell (Normal style of this file)</t></is></c><c r="B1" s="1" t="inlineStr"><is><t>bold</t></is></c></row>',
        '<row r="2"><c r="A2"><v>42</v></c></row>']
make_min_xlsx(p('gc_normal_font_arial10.xlsx'), rows, ('Arial', '10'))
make_min_xlsx(p('gc_normal_font_aptos.xlsx'), rows, ('Aptos Narrow', '11'))

# --- gc_implicit_intersection.xlsx -------------------------------------------------------------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Calc'
for r in range(1, 11):
    ws.cell(row=r, column=4, value=r * 100.0); ws.cell(row=r, column=6, value=r * 2.0); ws.cell(row=r, column=7, value=r)
wb.defined_names['Gesamtkosten'] = DefinedName('Gesamtkosten', attr_text='Calc!$D:$D')
wb.defined_names['Gesamt_EH'] = DefinedName('Gesamt_EH', attr_text='Calc!$F:$F')
ws['H9'] = '=Gesamtkosten/Gesamt_EH*G9'; ws['I9'] = '=D:D*2'; ws['J9'] = '=$D$1:$D$10+1'; ws['K9'] = '=SUM(D:D)'; ws['L9'] = '=Gesamtkosten'
ws['L10'] = 'blocker'   # a legacy formula never spills, so this cell must not cause #SPILL!
ws['A12'] = 'Legacy formulas (no dynamic-array cell metadata): Excel applies implicit intersection. Expected H9=450 I9=1800 J9=901 K9=5500 L9=900; nothing below row 9.'
wb.save(p('gc_implicit_intersection.xlsx'))

# --- gc_scoped_names.xlsx ----------------------------------------------------------------------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Calc'; ws2 = wb.create_sheet('Other')
for r in range(1, 11):
    ws.cell(row=r, column=4, value=r * 100.0); ws.cell(row=r, column=6, value=r * 2.0); ws.cell(row=r, column=7, value=r)
    ws2.cell(row=r, column=4, value=r * 1000.0); ws2.cell(row=r, column=6, value=r * 4.0); ws2.cell(row=r, column=7, value=r)
wb.defined_names['Gesamtkosten'] = DefinedName('Gesamtkosten', attr_text='Calc!$D:$D')
wb.defined_names['Gesamt_EH'] = DefinedName('Gesamt_EH', attr_text='Calc!$F:$F')
ws2.defined_names['Gesamtkosten'] = DefinedName('Gesamtkosten', attr_text='Other!$D:$D')
ws2.defined_names['Gesamt_EH'] = DefinedName('Gesamt_EH', attr_text='Other!$F:$F')
ws['H9'] = '=Gesamtkosten/Gesamt_EH*G9'; ws['I9'] = '=Gesamtkosten'
ws2['H9'] = '=Gesamtkosten/Gesamt_EH*G9'; ws2['I9'] = '=Gesamtkosten'
ws['A12'] = 'Global names Gesamtkosten/Gesamt_EH (Calc!$D:$D / $F:$F) + sheet-scoped duplicates on sheet Other. Expected Calc!H9=450 Calc!I9=900 Other!H9=2250 Other!I9=9000'
wb.save(p('gc_scoped_names.xlsx'))

# --- gc_coerce_de.xlsx -------------------------------------------------------------------------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Coerce'
for i, t in enumerate(['107,689', '1.330,00', '0,001', '1,234.5'], start=1):
    c = ws.cell(row=i, column=1, value=t); c.number_format = '@'
ws['B1'] = '=A1*1'; ws['B2'] = '=A2/2'; ws['B3'] = '=A3*1000'; ws['B4'] = '=A4*1'
ws['C1'] = '=VALUE(A1)'; ws['C2'] = '=VALUE(A2)'; ws['C3'] = '=A1+0'; ws['C4'] = '=SUM(A1:A3)'; ws['C5'] = '=--A2'
ws['D1'] = 'Expected in a de-DE Excel (decimal ",", thousands "."): B1=107.689 B2=665 B3=1 B4=#VALUE! C1=107.689 C2=1330 C3=107.689 C4=0 C5=1330'
wb.save(p('gc_coerce_de.xlsx'))

# --- gc_sort_collation.xlsx --------------------------------------------------------------------
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sort'
vals = ['x-1', 'x01', 'JO30-WE-1-', 'JO30-WE01-EGH', 'a-b', 'ab', 'AB', 'a c', 'a-c', 'ac', 'Ärger', 'Zebra', 'apfel', 'Apfel', '_x', 'x_', '1', '10', '2']
for i, v in enumerate(vals, start=1):
    c = ws.cell(row=i, column=1, value=v); c.number_format = '@'
# Excel's file syntax for the dynamic-array functions: SORT lives in the _xlws namespace, SORTBY in _xlfn.
# Written without the prefixes, GridCraft reads the file but Excel 365 refuses to open it (AC03, 13 Oct 2026).
ws['C1'] = '=_xlfn._xlws.SORT(A1:A19)'; ws['E1'] = '=_xlfn.SORTBY(A1:A19,A1:A19,-1)'
ws['G1'] = 'Excel 365 (de-DE) SORT asc: _x | 1 | 10 | 2 | a c | ab | AB | a-b | ac | a-c | apfel | Apfel | Ärger | JO30-WE01-EGH | JO30-WE-1- | x_ | x01 | x-1 | Zebra'
wb.save(p('gc_sort_collation.xlsx'))
print('ok')
