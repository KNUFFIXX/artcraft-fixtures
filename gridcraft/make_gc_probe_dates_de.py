#!/usr/bin/env python3
"""make_gc_probe_dates_de.py (DRACO-AC17, 27.10.2026) - probe for gridcraft#889: German date text read by date functions.
Sheet "Dates", no cached values: A1 '31.01.2021' (text), B1 =YEAR(A1), C1 =DATEVALUE(A1), D1 =A1+0;
A2 '16.10.2023 14:30' (text), B2 =DATEVALUE(A2), C2 =TIMEVALUE(A2), D2 =MONTH(A2).
Expected in a de-DE Excel (short date dd.MM.yyyy): B1=2021 C1=44227 D1=44227 B2=45215 C2=0.604166667 D2=10.
Usage: python3 make_gc_probe_dates_de.py <out.xlsx>   (needs openpyxl)"""
import sys
import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Dates'
ws['A1'] = '31.01.2021'; ws['B1'] = '=YEAR(A1)'; ws['C1'] = '=DATEVALUE(A1)'; ws['D1'] = '=A1+0'
ws['A2'] = '16.10.2023 14:30'; ws['B2'] = '=DATEVALUE(A2)'; ws['C2'] = '=TIMEVALUE(A2)'; ws['D2'] = '=MONTH(A2)'
ws['F1'] = 'Probe (no cached values). Expected in a de-DE Excel: B1=2021 C1=44227 D1=44227 B2=45215 C2=0.604166667 D2=10'
wb.save(sys.argv[1])
